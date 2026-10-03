"""Universal Document Parser Module for Knowledge Ingestion.

Supports:
- PDF (.pdf) via pypdf (with robust fallback)
- Microsoft Word (.docx) via python-docx & pure-python OOXML fallback
- Microsoft Excel (.xlsx) via openpyxl & pure-python OOXML fallback
- Microsoft PowerPoint (.pptx) via python-pptx & pure-python OOXML fallback
- Text & Markdown & CSV (.txt, .md, .csv, .tsv, .json) with auto-encoding detection (UTF-8, CP932)
"""

from __future__ import annotations

import csv
import io
import json
import logging
import os
import re
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)


@dataclass
class ExtractedSection:
    """A logical section extracted from a document (e.g. a page, a slide, a sheet)."""
    document_title: str
    locator: str
    text: str
    section_index: int = 0
    metadata: Dict[str, Any] = field(default_factory=dict)


@dataclass
class DocumentParseResult:
    """Overall outcome of parsing a document file."""
    filename: str
    file_type: str
    total_characters: int
    sections: List[ExtractedSection]
    metadata: Dict[str, Any] = field(default_factory=dict)


class UniversalDocumentParser:
    """Parses various binary and text document formats into normalized extracted sections."""

    @classmethod
    def parse_file(cls, filename: str, file_bytes: bytes) -> DocumentParseResult:
        """Parse file bytes based on filename extension."""
        ext = os.path.splitext(filename.lower())[1]

        if ext == ".pdf":
            return cls._parse_pdf(filename, file_bytes)
        elif ext == ".docx":
            return cls._parse_docx(filename, file_bytes)
        elif ext == ".xlsx":
            return cls._parse_xlsx(filename, file_bytes)
        elif ext == ".pptx":
            return cls._parse_pptx(filename, file_bytes)
        elif ext in [".txt", ".md", ".json"]:
            return cls._parse_text(filename, file_bytes, ext[1:])
        elif ext in [".csv", ".tsv"]:
            return cls._parse_csv(filename, file_bytes, delimiter="\t" if ext == ".tsv" else ",")
        else:
            # Fallback to plain text attempt
            return cls._parse_text(filename, file_bytes, "unknown")

    # ---------------------------------------------------------
    # PDF Parser
    # ---------------------------------------------------------
    @classmethod
    def _parse_pdf(cls, filename: str, file_bytes: bytes) -> DocumentParseResult:
        sections: List[ExtractedSection] = []
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                clean_text = cls._clean_text(page_text)
                if clean_text:
                    sections.append(ExtractedSection(
                        document_title=filename,
                        locator=f"ページ {idx + 1}",
                        text=clean_text,
                        section_index=idx + 1,
                        metadata={"page": idx + 1, "total_pages": len(reader.pages)},
                    ))
        except Exception as e:
            logger.warning(f"pypdf extraction error for '{filename}': {e}. Using raw text stream extraction.")
            # Fallback stream search
            raw_text = re.sub(rb"[\r\n]+", b"\n", file_bytes)
            # Try to grab literal strings inside parentheses or stream objects
            matches = re.findall(rb"\((.*?)\)", raw_text)
            extracted = " ".join([m.decode("latin1", errors="ignore") for m in matches if len(m) > 2])
            clean_text = cls._clean_text(extracted)
            if clean_text:
                sections.append(ExtractedSection(
                    document_title=filename,
                    locator="本文全体 (フォールバック抽出)",
                    text=clean_text,
                    section_index=1,
                ))

        total_chars = sum(len(s.text) for s in sections)
        return DocumentParseResult(
            filename=filename,
            file_type="PDF",
            total_characters=total_chars,
            sections=sections,
            metadata={"page_count": len(sections)},
        )

    # ---------------------------------------------------------
    # Word (.docx) Parser
    # ---------------------------------------------------------
    @classmethod
    def _parse_docx(cls, filename: str, file_bytes: bytes) -> DocumentParseResult:
        sections: List[ExtractedSection] = []
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            current_heading = "第1章 概要"
            current_paras: List[str] = []

            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                if p.style and ("Heading" in p.style.name or "見出し" in p.style.name):
                    if current_paras:
                        sections.append(ExtractedSection(
                            document_title=filename,
                            locator=current_heading,
                            text="\n".join(current_paras),
                            section_index=len(sections) + 1,
                        ))
                        current_paras = []
                    current_heading = text
                else:
                    current_paras.append(text)

            # Tables in Word
            for t_idx, table in enumerate(doc.tables):
                table_lines = []
                for row in table.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    table_lines.append(" | ".join(cells))
                if table_lines:
                    current_paras.append(f"【表 {t_idx + 1}】\n" + "\n".join(table_lines))

            if current_paras:
                sections.append(ExtractedSection(
                    document_title=filename,
                    locator=current_heading,
                    text="\n".join(current_paras),
                    section_index=len(sections) + 1,
                ))

        except Exception as e:
            logger.warning(f"python-docx parsing failed for '{filename}': {e}. Using pure OOXML XML parser.")
            sections = cls._parse_docx_pure_xml(filename, file_bytes)

        total_chars = sum(len(s.text) for s in sections)
        return DocumentParseResult(
            filename=filename,
            file_type="Word (DOCX)",
            total_characters=total_chars,
            sections=sections,
        )

    @classmethod
    def _parse_docx_pure_xml(cls, filename: str, file_bytes: bytes) -> List[ExtractedSection]:
        """Pure-Python OOXML word/document.xml parser fallback."""
        sections: List[ExtractedSection] = []
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                if "word/document.xml" in z.namelist():
                    xml_content = z.read("word/document.xml")
                    root = ET.fromstring(xml_content)
                    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                    paragraphs = []
                    for p in root.findall(".//w:p", ns):
                        texts = [t.text for t in p.findall(".//w:t", ns) if t.text]
                        if texts:
                            paragraphs.append("".join(texts))
                    if paragraphs:
                        sections.append(ExtractedSection(
                            document_title=filename,
                            locator="Word本文 (OOXML解析)",
                            text="\n".join(paragraphs),
                            section_index=1,
                        ))
        except Exception as err:
            logger.error(f"Pure OOXML docx parse failed: {err}")
        return sections

    # ---------------------------------------------------------
    # Excel (.xlsx) Parser
    # ---------------------------------------------------------
    @classmethod
    def _parse_xlsx(cls, filename: str, file_bytes: bytes) -> DocumentParseResult:
        sections: List[ExtractedSection] = []
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows_text = []
                for row in sheet.iter_rows(values_only=True):
                    # Filter empty rows
                    cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if cells:
                        rows_text.append(" | ".join(cells))

                if rows_text:
                    sections.append(ExtractedSection(
                        document_title=filename,
                        locator=f"シート: {sheet_name}",
                        text="\n".join(rows_text),
                        section_index=len(sections) + 1,
                        metadata={"sheet_name": sheet_name, "row_count": len(rows_text)},
                    ))
        except Exception as e:
            logger.warning(f"openpyxl failed for '{filename}': {e}. Using pure OOXML Excel parser.")
            sections = cls._parse_xlsx_pure_xml(filename, file_bytes)

        total_chars = sum(len(s.text) for s in sections)
        return DocumentParseResult(
            filename=filename,
            file_type="Excel (XLSX)",
            total_characters=total_chars,
            sections=sections,
            metadata={"sheet_count": len(sections)},
        )

    @classmethod
    def _parse_xlsx_pure_xml(cls, filename: str, file_bytes: bytes) -> List[ExtractedSection]:
        """Pure-Python OOXML sharedStrings and sheet xml parser fallback."""
        sections: List[ExtractedSection] = []
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                shared_strings = []
                if "xl/sharedStrings.xml" in z.namelist():
                    ss_root = ET.fromstring(z.read("xl/sharedStrings.xml"))
                    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    for si in ss_root.findall(".//x:si", ns):
                        t = si.find(".//x:t", ns)
                        shared_strings.append(t.text if t is not None and t.text else "")

                # Sheets
                sheet_files = [f for f in z.namelist() if f.startswith("xl/worksheets/sheet")]
                for idx, sf in enumerate(sheet_files):
                    sheet_root = ET.fromstring(z.read(sf))
                    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    lines = []
                    for row in sheet_root.findall(".//x:row", ns):
                        row_cells = []
                        for c in row.findall(".//x:c", ns):
                            v = c.find(".//x:v", ns)
                            if v is not None and v.text:
                                if c.get("t") == "s" and int(v.text) < len(shared_strings):
                                    row_cells.append(shared_strings[int(v.text)])
                                else:
                                    row_cells.append(v.text)
                        if row_cells:
                            lines.append(" | ".join(row_cells))
                    if lines:
                        sections.append(ExtractedSection(
                            document_title=filename,
                            locator=f"シート #{idx + 1} (OOXML)",
                            text="\n".join(lines),
                            section_index=idx + 1,
                        ))
        except Exception as err:
            logger.error(f"Pure OOXML xlsx parse failed: {err}")
        return sections

    # ---------------------------------------------------------
    # PowerPoint (.pptx) Parser
    # ---------------------------------------------------------
    @classmethod
    def _parse_pptx(cls, filename: str, file_bytes: bytes) -> DocumentParseResult:
        sections: List[ExtractedSection] = []
        try:
            import pptx
            prs = pptx.Presentation(io.BytesIO(file_bytes))
            for idx, slide in enumerate(prs.slides):
                slide_texts = []
                slide_title = ""
                if slide.shapes.title and slide.shapes.title.text:
                    slide_title = slide.shapes.title.text.strip()
                    slide_texts.append(f"【スライド見出し】{slide_title}")

                for shape in slide.shapes:
                    if shape.has_text_frame and shape != slide.shapes.title:
                        for paragraph in shape.text_frame.paragraphs:
                            t = paragraph.text.strip()
                            if t:
                                slide_texts.append(t)

                if slide_texts:
                    sections.append(ExtractedSection(
                        document_title=filename,
                        locator=f"スライド #{idx + 1} {f'({slide_title})' if slide_title else ''}".strip(),
                        text="\n".join(slide_texts),
                        section_index=idx + 1,
                        metadata={"slide": idx + 1, "title": slide_title},
                    ))
        except Exception as e:
            logger.warning(f"python-pptx failed for '{filename}': {e}. Using pure OOXML PowerPoint parser.")
            sections = cls._parse_pptx_pure_xml(filename, file_bytes)

        total_chars = sum(len(s.text) for s in sections)
        return DocumentParseResult(
            filename=filename,
            file_type="PowerPoint (PPTX)",
            total_characters=total_chars,
            sections=sections,
            metadata={"slide_count": len(sections)},
        )

    @classmethod
    def _parse_pptx_pure_xml(cls, filename: str, file_bytes: bytes) -> List[ExtractedSection]:
        """Pure-Python OOXML slide*.xml parser fallback."""
        sections: List[ExtractedSection] = []
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                slide_files = sorted([f for f in z.namelist() if f.startswith("ppt/slides/slide")])
                for idx, sf in enumerate(slide_files):
                    root = ET.fromstring(z.read(sf))
                    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                    texts = [t.text for t in root.findall(".//a:t", ns) if t.text and t.text.strip()]
                    if texts:
                        sections.append(ExtractedSection(
                            document_title=filename,
                            locator=f"スライド #{idx + 1} (OOXML)",
                            text="\n".join(texts),
                            section_index=idx + 1,
                        ))
        except Exception as err:
            logger.error(f"Pure OOXML pptx parse failed: {err}")
        return sections

    # ---------------------------------------------------------
    # Text / Markdown / CSV Parser
    # ---------------------------------------------------------
    @classmethod
    def _parse_text(cls, filename: str, file_bytes: bytes, file_type: str) -> DocumentParseResult:
        text = cls._decode_bytes(file_bytes)
        clean = cls._clean_text(text)
        sections = [
            ExtractedSection(
                document_title=filename,
                locator="本文全体",
                text=clean,
                section_index=1,
            )
        ]
        return DocumentParseResult(
            filename=filename,
            file_type=file_type.upper(),
            total_characters=len(clean),
            sections=sections,
        )

    @classmethod
    def _parse_csv(cls, filename: str, file_bytes: bytes, delimiter: str = ",") -> DocumentParseResult:
        text = cls._decode_bytes(file_bytes)
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        lines = []
        for row in reader:
            cleaned_row = [c.strip() for c in row if c.strip()]
            if cleaned_row:
                lines.append(" | ".join(cleaned_row))

        formatted_text = "\n".join(lines)
        sections = [
            ExtractedSection(
                document_title=filename,
                locator="データ一覧 (CSV/TSV)",
                text=formatted_text,
                section_index=1,
                metadata={"row_count": len(lines)},
            )
        ]
        return DocumentParseResult(
            filename=filename,
            file_type="CSV/TSV Table",
            total_characters=len(formatted_text),
            sections=sections,
            metadata={"row_count": len(lines)},
        )

    # ---------------------------------------------------------
    # Utilities
    # ---------------------------------------------------------
    @staticmethod
    def _decode_bytes(data: bytes) -> str:
        """Decode bytes with automatic charset detection (UTF-8, CP932, Shift_JIS, EUC-JP)."""
        encodings = ["utf-8", "utf-8-sig", "cp932", "shift_jis", "euc-jp", "latin1"]
        for enc in encodings:
            try:
                return data.decode(enc)
            except (UnicodeDecodeError, LookupError):
                continue
        return data.decode("utf-8", errors="replace")

    @staticmethod
    def _clean_text(text: str) -> str:
        """Normalize line breaks and trailing spaces."""
        text = re.sub(r"\r\n|\r", "\n", text)
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()
