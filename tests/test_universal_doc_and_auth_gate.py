"""Unit and Integration Tests for Universal Document Ingestion & Strict Auth Gate.

Covers:
- TEST-023: Strict Authentication Gate (401 Guard, Login, Session Management, Logout Invalidation)
- TEST-024: Universal Multi-Format Document Extraction (PDF, Word, Excel, PPTX, CSV, TXT)
"""

from __future__ import annotations

import io
import os
import tempfile
import unittest

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import pptx
except ImportError:
    pptx = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from pypdf import PdfWriter
except ImportError:
    PdfWriter = None

from src.core.engine import ProviderType
from src.rag.parser import UniversalDocumentParser
from src.web.app import KnowledgeWebApp


class TestUniversalDocAndAuthGate(unittest.TestCase):
    """Test suite for Strict Auth Gate and Universal Document Ingestion."""

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.chat_db = os.path.join(self.temp_dir.name, "test_chat.db")
        self.audit_db = os.path.join(self.temp_dir.name, "test_audit.db")

        self.app = KnowledgeWebApp(
            store_path=":memory:",
            default_provider=ProviderType.LITERT,
            gemini_api_key="AIzaSyTestSecretEnterpriseKey1234",
            gemini_model="gemini-3.5-flash-lite",
            chat_db_path=self.chat_db,
            audit_db_path=self.audit_db,
        )

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    # =========================================================
    # TEST-023: Strict Authentication Gate & Session Lifecycle
    # =========================================================
    def test_test_023_strict_auth_gate_and_session_lifecycle(self) -> None:
        """TEST-023: Verify unauthenticated requests are rejected, and logout invalidates session."""
        # 1. Unauthenticated requests to protected endpoints raise PermissionError
        with self.assertRaises(PermissionError):
            self.app.upload_document_file(
                filename="secret.docx",
                file_bytes=b"dummy",
                token=None,  # No token
            )

        with self.assertRaises(PermissionError):
            self.app.create_user(
                username="unauthorized_user",
                display_name="不正ユーザー",
                role="admin",
                department="無所属",
                token="invalid_dummy_token_123",
            )

        with self.assertRaises(PermissionError):
            self.app.set_retention_days(90, token=None)

        # 2. Login with valid credentials establishes session
        login_res = self.app.login(username="admin_user", password="password123")
        self.assertEqual(login_res["status"], "SUCCESS")
        self.assertIn("session_id", login_res)
        valid_token = login_res["session_id"]

        # 3. Operations succeed with valid session token
        ret_days = self.app.set_retention_days(90, token=valid_token)
        self.assertEqual(ret_days, 90)

        # 4. Logout invalidates session
        logout_res = self.app.logout(token=valid_token)
        self.assertEqual(logout_res["status"], "SUCCESS")

        # 5. Subsequent calls with the logged-out token are strictly rejected
        with self.assertRaises(PermissionError):
            self.app.set_retention_days(180, token=valid_token)

    # =========================================================
    # TEST-024: Universal Multi-Format Document Ingestion
    # =========================================================
    def test_test_024_text_and_csv_ingestion(self) -> None:
        """TEST-024: Parse and ingest TXT and CSV formats."""
        # 1. Text Ingestion
        txt_content = (
            "【全社規程】\n"
            "フレックスタイム制のコアタイムは11:00から15:00までと定めます。\n"
            "清算期間は毎月1日から末日までとします。"
        ).encode("utf-8")
        res_txt = UniversalDocumentParser.parse_file("勤務規程.txt", txt_content)
        self.assertEqual(res_txt.file_type, "TXT")
        self.assertTrue(res_txt.total_characters > 20)
        self.assertIn("コアタイム", res_txt.sections[0].text)

        # 2. CSV Ingestion
        csv_content = (
            "役職,役職手当,出張日当\n"
            "部長,80000,4000\n"
            "課長,50000,3000\n"
            "一般社員,0,2000"
        ).encode("utf-8")
        res_csv = UniversalDocumentParser.parse_file("役職手当一覧.csv", csv_content)
        self.assertEqual(res_csv.file_type, "CSV/TSV Table")
        self.assertIn("部長 | 80000 | 4000", res_csv.sections[0].text)

    def test_test_024_word_docx_ingestion(self) -> None:
        """TEST-024: Create and parse actual Microsoft Word (.docx) document."""
        if Document is None:
            self.skipTest("docx library not installed")
        doc = Document()
        doc.add_heading("第1章 福利厚生制度", level=1)
        doc.add_paragraph("社員の健康増進のため、フィットネスジム利用料の半額（月額上限5,000円）を補助します。")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "対象者"
        table.cell(0, 1).text = "補助率"
        table.cell(1, 0).text = "正社員・契約社員"
        table.cell(1, 1).text = "50%"

        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        # Parse with UniversalDocumentParser
        res = UniversalDocumentParser.parse_file("福利厚生ガイド.docx", docx_bytes)
        self.assertEqual(res.file_type, "Word (DOCX)")
        self.assertTrue(len(res.sections) >= 1)
        full_text = " ".join([s.text for s in res.sections])
        self.assertIn("フィットネスジム利用料", full_text)
        self.assertIn("月額上限5,000円", full_text)
        self.assertIn("正社員・契約社員", full_text)

    def test_test_024_excel_xlsx_ingestion(self) -> None:
        """TEST-024: Create and parse actual Microsoft Excel (.xlsx) workbook."""
        if openpyxl is None:
            self.skipTest("openpyxl library not installed")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "慶弔見舞金基準"
        ws.append(["事由", "勤続年数", "支給金額"])
        ws.append(["結婚祝金", "3年以上", "50,000円"])
        ws.append(["出産祝金", "全対象", "30,000円"])

        buf = io.BytesIO()
        wb.save(buf)
        xlsx_bytes = buf.getvalue()

        res = UniversalDocumentParser.parse_file("慶弔見舞金.xlsx", xlsx_bytes)
        self.assertEqual(res.file_type, "Excel (XLSX)")
        self.assertEqual(len(res.sections), 1)
        self.assertIn("慶弔見舞金基準", res.sections[0].locator)
        self.assertIn("結婚祝金 | 3年以上 | 50,000円", res.sections[0].text)

    def test_test_024_powerpoint_pptx_ingestion(self) -> None:
        """TEST-024: Create and parse actual Microsoft PowerPoint (.pptx) presentation."""
        if pptx is None:
            self.skipTest("pptx library not installed")
        prs = pptx.Presentation()
        slide_layout = prs.slide_layouts[1]  # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = "新入社員セキュリティ研修"
        tf = slide.shapes.placeholders[1].text_frame
        tf.text = "重要セキュリティ原則"
        p = tf.add_paragraph()
        p.text = "パスワードは12文字以上で大文字小文字記号を含めること。"

        buf = io.BytesIO()
        prs.save(buf)
        pptx_bytes = buf.getvalue()

        res = UniversalDocumentParser.parse_file("セキュリティ研修.pptx", pptx_bytes)
        self.assertEqual(res.file_type, "PowerPoint (PPTX)")
        self.assertEqual(len(res.sections), 1)
        self.assertIn("新入社員セキュリティ研修", res.sections[0].text)
        self.assertIn("パスワードは12文字以上", res.sections[0].text)

    def test_test_024_pdf_ingestion(self) -> None:
        """TEST-024: Create and parse actual PDF (.pdf) file."""
        if PdfWriter is None:
            self.skipTest("pypdf library not installed")
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)

        buf = io.BytesIO()
        writer.write(buf)
        pdf_bytes = buf.getvalue()

        res = UniversalDocumentParser.parse_file("白紙規程.pdf", pdf_bytes)
        self.assertEqual(res.file_type, "PDF")

    def test_test_024_end_to_end_file_upload_and_chat_retrieval(self) -> None:
        """TEST-024: End-to-end test: Ingest Word docx via API and retrieve answer via RAG chat."""
        if Document is None:
            self.skipTest("docx library not installed")
        # 1. Login as editor
        login = self.app.login("editor_user", "password123")
        editor_token = login["session_id"]

        # 2. Build Word docx with specific test fact
        doc = Document()
        doc.add_heading("特別休暇規定", level=1)
        doc.add_paragraph("ボランティア活動に参加する場合、年間最大5日間の特別有給休暇（ボランティア休暇）を取得できます。")
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        # 3. Upload file via upload_document_file API
        report = self.app.upload_document_file(
            filename="特別休暇規定_2026.docx",
            file_bytes=docx_bytes,
            token=editor_token,
            ip_address="192.168.1.150",
        )
        self.assertEqual(report["status"], "SUCCESS")
        self.assertEqual(report["file_type"], "Word (DOCX)")
        self.assertTrue(report["indexed_blocks"] >= 1)

        # 4. Query via Chat Assistant
        chat_res = self.app.chat("ボランティア休暇の日数は何日ですか？", token=editor_token)
        self.assertIn("answer", chat_res)
        self.assertTrue(len(chat_res["citation_links"]) > 0)
        # Verify retrieved contexts contains our uploaded document
        ctx_docs = [c["document_title"] for c in chat_res["retrieved_contexts"]]
        self.assertIn("特別休暇規定_2026.docx", ctx_docs)


if __name__ == "__main__":
    unittest.main()
